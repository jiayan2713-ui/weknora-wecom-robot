"""Prepare evidence views without changing source documents or following links."""
from __future__ import annotations

import csv
import io
import json
import os
import posixpath
import re
import subprocess
import signal
import zipfile
from contextlib import nullcontext

from .pipeline import ordered_bounded_map
from lxml import etree as ET
from pathlib import Path
from xml.sax.saxutils import escape

from .input_policy import unicode_evidence, native_sheet_units
from .common import WorkflowError, atomic_text, write_json, digest

VERSION = "prepare-xlsx-regions-v2"
SUPPORTED = {".txt", ".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".tsv", ".png", ".jpg", ".jpeg", ".webp"}


SKIPPED_MEDIA = {
    ".mp4", ".m4v", ".mov", ".avi", ".mkv", ".webm", ".wmv", ".flv", ".mpg", ".mpeg", ".3gp",
    ".m4a", ".mp3", ".wav", ".aac", ".flac", ".ogg", ".opus", ".wma", ".aiff", ".amr",
    ".wemta", ".wemtv", ".wemtc", ".wemtaidx", ".wemtvidx", ".m3u", ".m3u8", ".idx",
}


def text_units(source):
    if source.stat().st_size > 16 * 1024 * 1024:
        raise WorkflowError("CAPACITY_EXCEEDED", "TXT超过16MiB，请按行拆分")
    raw = source.read_bytes()
    encodings = ["utf-16"] if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else ["utf-8-sig", "gb18030"]
    for encoding in encodings:
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise WorkflowError("INVALID_DOCUMENT", "TXT无法按UTF-8、带BOM的UTF-16或GB18030解码")
    if "\x00" in text:
        raise WorkflowError("INVALID_DOCUMENT", "TXT含NUL字符，请检查编码或文件类型")
    units, chunk, size, start = [], [], 0, 1
    for number, line in enumerate(text.splitlines(keepends=True), 1):
        if len(line) > 12000:
            raise WorkflowError("CAPACITY_EXCEEDED", "TXT单行超过12000字符，请先按语义换行")
        if chunk and (size + len(line) > 12000 or len(chunk) >= 200):
            units.append({"label": f"原文件第{start}–{number-1}行", "text": "".join(chunk), "image": None})
            chunk, size, start = [], 0, number
        chunk.append(line)
        size += len(line)
    if chunk:
        units.append({"label": f"原文件第{start}–{number}行", "text": "".join(chunk), "image": None})
    return units


def run(command, timeout, *, env=None, network_denied=False):
    from .runtime import check_stop, bounded_timeout
    check_stop()
    timeout=bounded_timeout(timeout)
    if network_denied:
        guard = Path("/usr/bin/sandbox-exec")
        if not guard.exists():
            raise WorkflowError("CONVERSION_SAFETY_UNAVAILABLE", "缺少转换进程禁网络机制")
        command = [str(guard), "-p", "(version 1) (allow default) (deny network*)", *command]
    try:
        proc = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=env, start_new_session=True)
        try:
            stdout, stderr = proc.communicate(timeout=timeout)
        except subprocess.TimeoutExpired:
            os.killpg(proc.pid, signal.SIGKILL)
            proc.communicate()
            raise
    except subprocess.TimeoutExpired as exc:
        raise WorkflowError("CONVERSION_TIMEOUT", "转换超过时间预算", scope={"timeout_seconds": timeout}) from exc
    except OSError as exc:
        raise WorkflowError("CONVERSION_TOOL_MISSING", "转换工具不存在或无法启动") from exc
    if proc.returncode:
        # Do not copy arbitrary subprocess diagnostics, which may include document contents.
        code = "CONVERSION_SAFETY_UNAVAILABLE" if b"sandbox_apply" in stderr else "CONVERSION_FAILED"
        raise WorkflowError(code, f"转换进程失败（退出码{proc.returncode}）")


def office_profile(work):
    profile = work / "office-profile"
    user = profile / "user"
    user.mkdir(parents=True, exist_ok=True)
    atomic_text(user / "registrymodifications.xcu", '''<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry">
<item oor:path="/org.openoffice.Office.Common/Security/Scripting">
 <prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop>
 <prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop>
</item>
<item oor:path="/org.openoffice.Office.Update/Update">
 <prop oor:name="Enabled" oor:op="fuse"><value>false</value></prop>
</item>
</oor:items>''')
    fonts = [Path("/System/Library/Fonts"), Path("/Library/Fonts")]
    conf = work / "fonts.conf"
    atomic_text(conf, '<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig>' +
                ''.join(f'<dir>{escape(str(p))}</dir>' for p in fonts if p.exists()) +
                f'<cachedir>{escape(str(work / "font-cache"))}</cachedir></fontconfig>')
    env = {k: v for k, v in os.environ.items() if not any(x in k.upper() for x in ["KEY", "TOKEN", "SECRET", "PASSWORD"])}
    env["FONTCONFIG_FILE"] = str(conf)
    return profile, env


def trim_full_width_style_tail(tree):
    """Remove only an XFD-wide empty style tail, retaining content and merged ranges."""
    from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries, get_column_letter
    dimension = next((n for n in tree if n.tag.rsplit("}", 1)[-1] == "dimension"), None)
    if dimension is None:
        return None
    _, _, max_col, max_row = range_boundaries(dimension.get("ref", "A1"))
    if max_col != 16384 or max_col * max_row <= 300000:
        return None
    blocked = {"drawing", "legacyDrawing", "hyperlinks", "dataValidations", "conditionalFormatting", "tableParts", "extLst"}
    if any(n.tag.rsplit("}", 1)[-1] in blocked for n in tree.iter()):
        return None
    cells = [n for n in tree.iter() if n.tag.rsplit("}", 1)[-1] == "c"]
    content = [n for n in cells if len(n)]
    if not content:
        return None
    bounds = [coordinate_to_tuple(n.get("r")) for n in content]
    end_row = max(r for r, c in bounds)
    end_col = max(c for r, c in bounds)
    for node in tree.iter():
        if node.tag.rsplit("}", 1)[-1] == "mergeCell":
            _, _, c, r = range_boundaries(node.get("ref"))
            end_col, end_row = max(end_col, c), max(end_row, r)
    if end_col >= 16384 or end_col * max(end_row, max_row) > 300000:
        return None
    removed = 0
    for node in cells:
        row, col = coordinate_to_tuple(node.get("r"))
        if col > end_col and not len(node):
            node.getparent().remove(node)
            removed += 1
    if not removed:
        return None
    old = dimension.get("ref")
    new = f"A1:{get_column_letter(end_col)}{max(end_row, max_row)}"
    dimension.set("ref", new)
    return f"转换副本移除延伸至XFD的{removed}个右侧纯格式空单元格，范围{old}改为{new}；保留原件、内容与合并区域，版式待审核。"


def embedded_image_fallbacks(archive):
    """Detach redundant external image links only when the same blip embeds an image."""
    changes, notes = {}, []
    rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    parser = ET.XMLParser(resolve_entities=False, no_network=True)
    for name in archive.namelist():
        if not (name.startswith("xl/drawings/_rels/") and name.endswith(".xml.rels")):
            continue
        raw = archive.read(name)
        if b"<!DOCTYPE" in raw or b"<!ENTITY" in raw:
            continue  # The normal safety scan rejects these.
        tree = ET.fromstring(raw, parser)
        owner = name.replace("/_rels/", "/")[:-5]
        owner_raw = archive.read(owner)
        if b"<!DOCTYPE" in owner_raw or b"<!ENTITY" in owner_raw:
            continue
        drawing = ET.fromstring(owner_raw, parser)
        relationships = {r.get("Id"): r for r in tree}
        changed = False
        for rel in list(tree):
            if rel.get("TargetMode") != "External" or not rel.get("Type", "").endswith("/image"):
                continue
            references = [n for n in drawing.iter() if rel.get("Id") in n.attrib.values()]
            if not references:
                continue
            valid = True
            for node in references:
                embedded = relationships.get(node.get(rel_ns + "embed"))
                if (node.tag.rsplit("}", 1)[-1] != "blip" or node.get(rel_ns + "link") != rel.get("Id")
                        or embedded is None or embedded.get("TargetMode") == "External"
                        or not embedded.get("Type", "").endswith("/image")):
                    valid = False
                    break
                target = posixpath.normpath(posixpath.join(posixpath.dirname(owner), embedded.get("Target", "")))
                if not target.startswith("xl/media/") or target not in archive.namelist() or archive.getinfo(target).file_size == 0:
                    valid = False
                    break
            if not valid:
                continue
            for node in references:
                del node.attrib[rel_ns + "link"]
            tree.remove(rel)
            notes.append({"label": "转换处理记录（非原文）", "text":
                          "图片同时存在内嵌副本和外部引用；转换副本仅使用内嵌图片，未访问外部目标。原始关系保留在归档原件：" + name + " " + str(rel.get("Id"))})
            changed = True
        if changed:
            changes[name] = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
            changes[owner] = ET.tostring(drawing, encoding="utf-8", xml_declaration=True)
    return changes, notes


def page_image_units(image, label, text):
    """Keep full resolution and source coordinates when an image exceeds the payload cap."""
    from PIL import Image
    if image.stat().st_size <= 8 * 1024 * 1024:
        return [{"label": label, "text": text, "image": image.name}]
    units = []
    with Image.open(image) as source:
        source = source.convert("RGB")
        def partition(box):
            tile = image.with_name(image.stem + "-region-" + "-".join(map(str, box)) + ".png")
            source.crop(box).save(tile)
            if tile.stat().st_size > 8 * 1024 * 1024:
                tile.unlink()
                x0, y0, x1, y1 = box
                if max(x1-x0, y1-y0) < 128:
                    raise WorkflowError("CAPACITY_EXCEEDED", "最小页面分区仍超限")
                if x1-x0 >= y1-y0:
                    mid = (x0+x1)//2
                    partition((x0, y0, mid+16, y1)); partition((mid-16, y0, x1, y1))
                else:
                    mid = (y0+y1)//2
                    partition((x0, y0, x1, mid+16)); partition((x0, mid-16, x1, y1))
                return
            units.append({"label": label + "；像素区域" + str(box), "text": text +
                          "\n本图为原页分区，相邻分区有32像素重叠；上方文字是整页提取参考，仅转写本区域可见内容。",
                          "image": tile.name,"region":{"kind":"pixel_tile","box":list(box),"page_size":[source.width,source.height],"recovery_depth":0}})
        partition((0, 0, source.width, source.height))
    return units


def safe_office_copy(source, work):
    """Reject active/external loads; retain ordinary hyperlinks as source data."""
    target = work / ("render-input" + source.suffix.lower())
    notes, links = [], []
    try:
        archive = zipfile.ZipFile(source)
    except zipfile.BadZipFile as exc:
        raise WorkflowError("INVALID_DOCUMENT", "Office文件损坏或加密") from exc
    with archive as incoming:
        entries = incoming.infolist()
        if len(entries) > 20000 or sum(x.file_size for x in entries) > 1024 ** 3:
            raise WorkflowError("CAPACITY_EXCEEDED", "Office解压内容超过本地处理容量")
        if sum(x.file_size for x in entries if x.filename.endswith((".xml", ".rels"))) > 64 * 1024 * 1024:
            raise WorkflowError("CAPACITY_EXCEEDED", "Office结构文本超过64MiB，请拆分")
        if len({x.filename for x in entries}) != len(entries):
            raise WorkflowError("INVALID_DOCUMENT", "Office压缩包有重复条目")
        changes, fallback_notes = embedded_image_fallbacks(incoming) if source.suffix.lower() == ".xlsx" else ({}, [])
        notes.extend(fallback_notes)
        removed = set()
        if source.suffix.lower() == '.xlsx':
            from .wps_checkboxes import freeze
            frozen, removed, control_notes = freeze(incoming)
            changes.update(frozen)
            notes.extend(control_notes)
        for entry in entries:
            if entry.filename in removed:
                continue
            name = entry.filename
            if any(x in name.lower() for x in ["vbaproject", "activex", "/embeddings/", "/externallinks/", "connections.xml"]):
                raise WorkflowError("UNSAFE_OFFICE", "文件含宏、主动对象或外部数据连接；未执行转换", scope=name)
            if not name.endswith((".xml", ".rels")):
                continue
            original_raw = incoming.read(name)
            raw = changes.get(name, original_raw)
            if b"<!DOCTYPE" in raw or b"<!ENTITY" in raw:
                raise WorkflowError("UNSAFE_OFFICE", "不接受含自定义实体的Office XML")
            if name.endswith(".rels"):
                raw = raw.replace(b'xmlns="http://schemas.openxmlformats.org/officeDocument/2006/relationships"',
                                  b'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"')
            try:
                tree = ET.fromstring(raw, parser=ET.XMLParser(resolve_entities=False, no_network=True, remove_comments=True))
            except ET.XMLSyntaxError as exc:
                raise WorkflowError("INVALID_DOCUMENT", "Office XML无法读取", scope=name) from exc
            for node in tree.iter():
                local = node.tag.rsplit("}", 1)[-1]
                if local in {"oleObject", "altChunk", "control", "webVideoPr"}:
                    raise WorkflowError("UNSAFE_OFFICE", "文件含需单独处理的主动内容", scope=local)
                if name.endswith(".rels") and node.get("TargetMode") == "External":
                    if not node.get("Type", "").endswith("/hyperlink"):
                        raise WorkflowError("UNSAFE_OFFICE", "文件含外部模板、图像或数据引用", scope=name)
                    links.append(node.get("Target", ""))
                if local in {"f", "instrText"} and re.search(r"(?i)\b(IMAGE|WEBSERVICE|DDE|DDEAUTO|RTD|INCLUDEPICTURE|INCLUDETEXT|LINK)\b", node.text or ""):
                    raise WorkflowError("UNSAFE_OFFICE", "文件含外部求值字段", scope=name)
            if name.startswith("word/") and re.search(r"(?:comments|footnotes|endnotes)\.xml$", name):
                for node in tree:
                    text = "".join(x.text or "" for x in node.iter() if x.tag.rsplit("}", 1)[-1] in {"t", "delText"})
                    if text.strip():
                        notes.append({"label": f"原文附属内容 {name} {list(node.attrib.values())}", "text": text})
            if name == "word/document.xml":
                for n, paragraph in enumerate((x for x in tree.iter() if x.tag.rsplit("}", 1)[-1] == "p"), 1):
                    for revision in paragraph:
                        kind = revision.tag.rsplit("}", 1)[-1]
                        if kind in {"ins", "del", "moveFrom", "moveTo"}:
                            text = "".join(x.text or "" for x in revision.iter() if x.tag.rsplit("}", 1)[-1] in {"t", "delText"})
                            notes.append({"label": f"原文修订 段落{n} {kind}", "text": json.dumps({"attributes": revision.attrib, "text": text}, ensure_ascii=False)})
            if source.suffix.lower() == ".xlsx" and (name == 'xl/workbook.xml' or name.startswith('xl/worksheets/')):
                for node in tree.iter():
                    if node.tag.rsplit("}", 1)[-1] == "sheet":
                        node.set("state", "visible")
                    if node.tag.rsplit("}", 1)[-1] in {"row", "col"} and node.get("hidden"):
                        node.set("hidden", "0")
                for node in list(tree.iter()):
                    for child in list(node):
                        if child.tag.rsplit("}", 1)[-1] == "definedName" and child.get("name") == "_xlnm.Print_Area":
                            node.remove(child)
                if name.startswith('xl/worksheets/') and not any('comments' in x.filename.lower() for x in entries):
                    trimming = trim_full_width_style_tail(tree)
                    if trimming:
                        notes.append({"label": "转换处理记录（非原文）", "text": name + "：" + trimming})
                raw = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
            if raw != original_raw:
                changes[name] = raw
        with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as output:
            for entry in entries:
                if entry.filename in removed:
                    continue
                output.writestr(entry, changes[entry.filename] if entry.filename in changes else incoming.read(entry.filename))
    return target, notes, list(dict.fromkeys(links))


def pdf_annotation_links(page):
    """Read link metadata only; malformed references become explicit review notes."""
    from pypdf.generic import IndirectObject, NullObject
    from pypdf.errors import PdfReadError
    links, issues = [], []

    def resolve(value):
        seen = set()
        while isinstance(value, IndirectObject):
            key = (id(value.pdf), value.idnum, value.generation)
            if key in seen or len(seen) >= 16:
                raise ValueError("批注引用循环或过深")
            seen.add(key)
            value = value.get_object()
        return value

    try:
        raw = page.get("/Annots")
        annots = resolve(raw)
        if annots is None or isinstance(annots, NullObject):
            if isinstance(raw, IndirectObject):
                issues.append("批注数组引用失效")
            return links, issues
        if not isinstance(annots, (list, tuple)):
            return links, ["批注数组类型异常"]
        for number, ref in enumerate(annots, 1):
            try:
                annotation = resolve(ref)
                if not isinstance(annotation, dict):
                    issues.append(f"批注{number}引用失效或类型异常")
                    continue
                if "/A" not in annotation:
                    continue
                raw_action = annotation.get("/A")
                action = resolve(raw_action)
                if action is None or isinstance(action, NullObject):
                    if isinstance(raw_action, IndirectObject):
                        issues.append(f"批注{number}动作引用失效")
                    continue
                if not isinstance(action, dict):
                    issues.append(f"批注{number}动作类型异常")
                    continue
                if "/URI" not in action:
                    continue
                uri = resolve(action.get("/URI"))
                if isinstance(uri, str):
                    links.append(uri)
                elif isinstance(uri, bytes):
                    links.append(uri.decode("latin-1"))
                else:
                    issues.append(f"批注{number}链接引用失效或类型异常")
            except (PdfReadError, ValueError, TypeError, KeyError, RecursionError) as error:
                issues.append(f"批注{number}无法读取（{type(error).__name__}）")
    except (PdfReadError, ValueError, TypeError, KeyError, RecursionError) as error:
        issues.append(f"批注数组无法读取（{type(error).__name__}）")
    return links, issues


def prepare(source, work, config, render_semaphore=None):
    from PIL import Image, ImageOps
    from pypdf import PdfReader
    source, work = Path(source), Path(work)
    work.mkdir(parents=True, exist_ok=True)
    ext = source.suffix.lower()
    if ext not in SUPPORTED:
        raise WorkflowError("UNSUPPORTED_FORMAT", "当前不支持此格式", scope=ext)
    units, native = [], {"links": [], "sheets": [], "slides": []}
    if ext == ".txt":
        units = text_units(source)
    elif ext in {".csv", ".tsv"}:
        if source.stat().st_size > 16 * 1024 * 1024:
            raise WorkflowError("CAPACITY_EXCEEDED", "分隔文件超过16MiB，请按记录边界拆分")
        raw = source.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("gb18030")
        try:
            rows = list(csv.reader(io.StringIO(text), delimiter="\t" if ext == ".tsv" else ",", strict=True))
        except csv.Error as exc:
            raise WorkflowError("INVALID_DOCUMENT", "分隔文件记录边界无效") from exc
        for i in range(0, len(rows), 40):
            units.append({"label": f"记录{i + 1}–{min(i + 40, len(rows))}", "text": json.dumps(rows[i:i+40], ensure_ascii=False), "image": None})
    elif ext in {".png", ".jpg", ".jpeg", ".webp"}:
        target = work / "page-001.png"
        with Image.open(source) as image:
            if image.width * image.height > 40000000:
                raise WorkflowError("CAPACITY_EXCEEDED", "图片超过4000万像素，请先分区")
            image = ImageOps.exif_transpose(image).convert("RGB")
            image.thumbnail((3600, 3600))
            image.save(target)
        if target.stat().st_size > 8 * 1024 * 1024:
            raise WorkflowError("CAPACITY_EXCEEDED", "准备后的图像超过8MiB，请先分区")
        units.append({"label": "图片1", "text": "", "image": target.name})
    else:
        notes, render_input, pdf = [], source, source
        if ext != ".pdf":
            try:
                render_input, notes, native["links"] = safe_office_copy(source, work)
            except WorkflowError as error:
                if error.code != 'UNSAFE_OFFICE' or not config.get('office_static_fallback', False):
                    raise
                from .office_static import prepare_static
                render_input, notes = prepare_static(source, work / 'static')
                render_input, extra_notes, native['links'] = safe_office_copy(render_input, work)
                notes.extend(extra_notes)
            if ext == ".xlsx":
                import openpyxl
                with zipfile.ZipFile(source) as z:
                    worksheet_bytes = sum(info.file_size for info in z.infolist() if info.filename.startswith('xl/worksheets/'))
                if worksheet_bytes > 32 * 1024 * 1024:
                    raise WorkflowError("CAPACITY_EXCEEDED", "工作表XML过大，请拆分")
                with open(render_input, "rb") as bounded_input:
                    bounded_book = openpyxl.load_workbook(bounded_input, data_only=False)
                    bounds = {s.title: (s.max_row, s.max_column, s.calculate_dimension()) for s in bounded_book}
                    bounded_book.close()
                formulas = openpyxl.load_workbook(source, data_only=False)
                cached = openpyxl.load_workbook(source, data_only=True)
                try:
                    for sheet in formulas:
                        max_row, max_column, effective_range = bounds[sheet.title]
                        if max_row * max_column > 300000:
                            raise WorkflowError("CAPACITY_EXCEEDED", "工作表区域过大，需拆分")
                        cells = [{"coordinate": c.coordinate, "value": str(c.value), "type": c.data_type,
                                  "format": c.number_format, "cached": cached[sheet.title][c.coordinate].value,
                                  "comment": c.comment.text if c.comment else None}
                                 for row in sheet.iter_rows(max_row=max_row, max_col=max_column) for c in row if c.value is not None]
                        native["sheets"].append({"name": sheet.title, "state": sheet.sheet_state, "range": effective_range,
                                                 "merges": [str(x) for x in sheet.merged_cells], "cells": cells})
                finally:
                    formulas.close()
                    cached.close()
            if ext == ".pptx":
                from pptx import Presentation
                for slide in Presentation(render_input).slides:
                    notes_frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
                    native["slides"].append({"notes": notes_frame.text if notes_frame is not None else "",
                                             "hidden": slide._element.get("show") == "0"})
            profile, env = office_profile(work)
            export_filter = "pdf"
            if ext == ".xlsx":
                export_filter = 'pdf:calc_pdf_Export:{"SinglePageSheets":{"type":"boolean","value":"true"}}'
            if ext == ".pptx":
                export_filter = 'pdf:impress_pdf_Export:{"ExportHiddenSlides":{"type":"boolean","value":"true"}}'
            command = [config["soffice"], "-env:UserInstallation=" + profile.as_uri(), "--headless", "--convert-to", export_filter,
                       "--outdir", str(work), str(render_input)]
            try:
                run(command, config["conversion_timeout_seconds"], env=env, network_denied=True)
            except WorkflowError as error:
                extended = config.get('conversion_retry_timeout_seconds', 300)
                if error.code != 'CONVERSION_TIMEOUT' or extended <= config['conversion_timeout_seconds']:
                    raise
                run(command, extended, env=env, network_denied=True)
                notes.append({'label': '转换处理记录（非原文）', 'text': '首次转换超时；一次有界延长预算后完成，需核对页数及对象。'})
            pdf = work / (render_input.stem + ".pdf")
        try:
            reader = PdfReader(pdf)
            if reader.is_encrypted and not reader.decrypt(""):
                raise WorkflowError("ENCRYPTED_DOCUMENT", "文件加密，未处理")
            count = len(reader.pages)
        except WorkflowError:
            raise
        except Exception as exc:
            raise WorkflowError("INVALID_DOCUMENT", "无法读取PDF或转换结果") from exc
        if not 0 < count <= 500:
            raise WorkflowError("CAPACITY_EXCEEDED", "页数为空或超过500页，需拆分")
        if ext == ".xlsx" and count != len(native["sheets"]):
            raise WorkflowError("LOCATION_MISMATCH", "工作表与渲染页不能一一对应")
        if ext == ".pptx" and count != len(native["slides"]):
            raise WorkflowError("LOCATION_MISMATCH", "幻灯片与渲染页不能一一对应")
        page_specs = []
        for i, page in enumerate(reader.pages):
            from .runtime import check_stop
            check_stop()
            label = f"PDF第{i+1}页" if ext == ".pdf" else f"本次固定渲染版第{i+1}页"
            text = page.extract_text() or ""
            text, unicode_note = unicode_evidence(text, i+1, work)
            if unicode_note:
                notes.append(unicode_note)
            if ext == ".xlsx":
                sheet = native["sheets"][i]
                label = f'工作表{sheet["name"]}！{sheet["range"]}'
                text += "\n原始值、公式及缓存（不得以渲染器重算代替）：\n" + json.dumps(sheet, ensure_ascii=False, default=str)
            if ext == ".pptx":
                label = f"幻灯片{i+1}"
                text += "\n原文演讲者备注：\n" + native["slides"][i]["notes"]
            links, annotation_issues = pdf_annotation_links(page)
            for uri in links:
                text += "\n原文链接（仅数据）：" + uri
            if annotation_issues:
                notes.append({"label": "转换处理记录（非原文）", "text":
                              f"{label}：" + "；".join(annotation_issues) +
                              "。已保留原件及页面，未补写或访问链接，需人工核对批注。"})
            native.setdefault('pages', []).append({'page': i+1, 'label': label, 'text': text})
            page_specs.append((i, label, text))
        def render_page(spec):
            i, label, text = spec
            image = work / f"page-{i+1:03}.png"
            with (render_semaphore if render_semaphore is not None else nullcontext()):
                if len(text.split('原始值、公式及缓存（不得以渲染器重算代替）：',1)[0]) > max(16000,config.get('max_unit_text_chars',16000)):
                    from .xlsx_regions import render_regions
                    return render_regions(pdf,i,label,work,config,page=reader.pages[i],text=text)
                run([config["pdftoppm"], "-f", str(i+1), "-l", str(i+1), "-r", "180", "-scale-to", "3600",
                     "-singlefile", "-png", str(pdf), str(image.with_suffix(""))], config["render_timeout_seconds"])
                page_units = page_image_units(image, label, text)
            if len(page_units) == 1 and ext in ('.pdf','.docx','.pptx'):
                page_units[0]['native_text'] = text
            for part in page_units:
                part['source_location'] = {'page': i+1, 'source_format': ext}
            return page_units
        native_supplements = set()
        for page_index, page_units in enumerate(ordered_bounded_map(render_page, page_specs, config.get("render_workers", 1))):
            for unit in page_units:
                if ext == '.xlsx' and unit.get('region'):
                    native_supplements.add(page_index)
                if ext == '.xlsx' and len(unit['text']) > config.get('max_unit_text_chars', 16000):
                    # Keep a visual unit, and preserve the oversized cell data separately.
                    native_supplements.add(page_index)
                    prefix = unit['text'].split('原始值、公式及缓存（不得以渲染器重算代替）：', 1)[0]
                    notice = '\n本工作表的完整原始单元格、公式、缓存另附原生记录；本图仅转写可见内容。'
                    unit['text'] = (prefix if len(prefix)+len(notice) <= config.get('max_unit_text_chars', 16000) else '') + notice
                units.append(unit)
        if ext == '.xlsx':
            for sheet_index, sheet in enumerate(native['sheets']):
                if sheet_index in native_supplements:
                    units.extend(native_sheet_units(sheet, config.get('max_unit_text_chars', 16000)))
        if ext in ('.pdf','.docx','.pptx'):
            for page in native.get('pages',[]):
                budget=config.get('max_unit_text_chars',16000)
                if len(page['text']) > budget:
                    for offset in range(0,len(page['text']),budget):
                        chunk=page['text'][offset:offset+budget]
                        units.append({'label': page['label']+'；原生文字片段'+str(offset),
                            'text': chunk, 'native_body': chunk, 'method': 'native_text', 'image': None,
                            'source_location': {'page':page['page'],'characters':[offset,offset+len(chunk)]}})
        units.extend({**note, "image": None} for note in notes)
    if not units:
        units = [{"label": "空文件", "text": "", "image": None}]
    for i, unit in enumerate(units):
        unit["id"] = f"u{i+1:04}"
    # Native datetime values are evidence text, not workflow state types.
    native = json.loads(json.dumps(native, ensure_ascii=False, default=str))
    write_json(work / "prepared.json", {"source_hash": digest(source), "version": VERSION, "units": units, "native": native})
    return units
