"""Deterministic XLSX projection and roundtrip checks, without model calls."""
import io
import json
import math

from .common import WorkflowError, atomic_bytes
from .evaluation import LEGACY_XLSX_FIELDS, XLSX_FIELDS, fact_quotes, row_values


def verify_workbook(path, candidate, facts, generation, table, *, columns=None):
    """Recover a completely written workbook whose pointer commit was interrupted."""
    import openpyxl
    try:
        book = openpyxl.load_workbook(path, data_only=False)
        try:
            if book.sheetnames != ['评测集', '原始证据链接', '说明']:
                raise ValueError('sheets')
            main, links = book['评测集'], book['原始证据链接']
            fields = XLSX_FIELDS if columns is None else columns
            if fields not in (LEGACY_XLSX_FIELDS, XLSX_FIELDS) or [c.value for c in main[1]] != fields or main.max_row != len(candidate)+1:
                raise ValueError('dimensions')
            quotes = fact_quotes(facts) if fields == XLSX_FIELDS else None
            index, count = 2, 0
            for n,item in enumerate(candidate,2):
                if [c.value for c in main[n]] != row_values(item,table,quotes):
                    raise ValueError('candidate')
                expected_main = table[item['evidence'][0]]['url'] if len(item['evidence']) == 1 else None
                actual_main = main.cell(n,5).hyperlink
                if (actual_main.target if actual_main else None) != expected_main:
                    raise ValueError('main hyperlink')
                for key in item['evidence']:
                    source=table[key]
                    expected=[item['id'],source['original_name'],source['location'],source['url'] or '未提供来源URL']
                    if [links.cell(index,c).value for c in range(1,5)] != expected:
                        raise ValueError('source')
                    if source['url']:
                        link=links.cell(index,4).hyperlink
                        if not link or link.target != source['url']:
                            raise ValueError('hyperlink')
                        count+=1
                    index+=1
            fact_header=index+2
            index=fact_header+1
            for record in facts:
                for source in record['sources']:
                    expected=[record['question_id'],str(record['fact_index']),record['fact'],source['evidence'],source['quote']]
                    if [links.cell(index,c).value for c in range(1,6)] != expected:
                        raise ValueError('facts')
                    index+=1
            if links.max_row != index-1 or any(c.data_type=='f' for sheet in book for row in sheet for c in row):
                raise ValueError('extra content')
            if book['说明']['B3'].value != generation['batch_id'] or book['说明']['B4'].value != generation['id']:
                raise ValueError('identity')
            return {'question_count':len(candidate),'columns':fields,'fact_count':len(facts),
                    'hyperlinks_verified':count,'roundtrip':'passed','status':'待业务人工审核'}
        finally:
            book.close()
    except WorkflowError:
        raise
    except Exception as exc:
        raise WorkflowError('EXPORT_MISMATCH', '未登记表格不能通过完整核验，已保留文件', scope=type(exc).__name__) from exc


def export_workbook(path, candidate, facts, generation, table, *, columns=None):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    book = openpyxl.Workbook()
    check = None
    try:
        main = book.active
        main.title = "评测集"
        links = book.create_sheet("原始证据链接")
        notes = book.create_sheet("说明")
        fields = XLSX_FIELDS if columns is None else columns
        if fields not in (LEGACY_XLSX_FIELDS, XLSX_FIELDS):
            raise WorkflowError('EXPORT_SCHEMA', '不支持的导出列契约')
        quotes = fact_quotes(facts) if fields == XLSX_FIELDS else None
        main.append(fields)
        links.append(["id", "原始文件名", "对应位置", "url"])
        link_targets = {}
        for number, item in enumerate(candidate, 2):
            main.append(row_values(item, table, quotes))
            for key in item["evidence"]:
                source = table[key]
                links.append([item["id"], source["original_name"], source["location"], source["url"] or "未提供来源URL"])
                if source["url"]:
                    address = f"D{links.max_row}"
                    links[address].hyperlink = source["url"]
                    link_targets[address] = source["url"]
            if len(item["evidence"]) == 1 and table[item["evidence"][0]]["url"]:
                main.cell(number, 5).hyperlink = table[item["evidence"][0]]["url"]
        fact_header = links.max_row + 3
        for i, value in enumerate(["id", "fact_index", "fact", "evidence", "quote"], 1):
            links.cell(fact_header, i, value)
        fact_rows = []
        for record in facts:
            for source in record["sources"]:
                row = [record["question_id"], str(record["fact_index"]), record["fact"], source["evidence"], source["quote"]]
                fact_rows.append(row)
                links.append(row)
        notes.append(["项目", "内容"])
        for row in [
            ["状态", "待业务人工审核"], ["批次", generation["batch_id"]], ["生成版本", generation["id"]],
            ["候选格式", generation["candidate_schema"]], ["事实核对格式", generation["fact_checks_schema"]],
            ["审核版本", generation["review"]["sha256"]], ["提示词版本指纹", generation["prompt"]["sha256"]],
            ["出题模型", generation["codex_model"]], ["程序检查", "字段、配比、引用、摘录存在性和表格往返；不证明语义正确"],
            ["人工审核", "确认事实原文依据、必要性、参考答案与推理正确性；库外/模糊题按其分类规则审核"],
            ["数组展示", "evidence为JSON文本；required_facts按编号展示，完整原值与来源见原始证据链接的事实明细"],
            ["审核正文摘录", "从fact_checks原样汇总；同一证据与摘录合并展示对应事实序号，完整关联仍保留于事实明细。摘录来自审核Markdown，不代表已逐字核验原始文件。"],
        ]:
            notes.append(row)
        for sheet in book:
            sheet.freeze_panes = "B2" if sheet == main else "A2"
            sheet.sheet_view.showGridLines = False
            for row in sheet:
                for cell in row:
                    if cell.value is not None:
                        cell.value = str(cell.value)
                        cell.data_type = "s"
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
                        cell.font = Font(name="Arial", size=11)
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="203B55")
                cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            sheet.row_dimensions[1].height = 26
        sizes = {"id":14,"question":48,"answer":82,"evidence":56,"original_evidence":62,
                 "required_facts":65,"审核正文摘录":100,"category":16,"推理步数":13}
        widths = {chr(65+i):sizes[field] for i,field in enumerate(fields)}
        for col, width in widths.items():
            main.column_dimensions[col].width = width
        for row in main.iter_rows(min_row=2):
            lines = max(sum(max(1, math.ceil(len(part)/(widths[c.column_letter]/1.7))) for part in str(c.value or "").split("\n")) for c in row)
            main.row_dimensions[row[0].row].height = min(409, max(45, lines * 15 + 10))
        main.auto_filter.ref = main.dimensions
        for col, width in {"A": 14, "B": 42, "C": 66, "D": 70, "E": 80}.items():
            links.column_dimensions[col].width = width
        notes.column_dimensions["A"].width = 24
        notes.column_dimensions["B"].width = 100
        for cell in links[fact_header]:
            cell.fill = PatternFill("solid", fgColor="203B55")
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        payload = io.BytesIO()
        book.save(payload)
        data = payload.getvalue()
        # Reopen serialized bytes before publishing an immutable workbook.
        check = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
        try:
            actual = check["评测集"]
            if [c.value for c in actual[1]] != fields or actual.max_row != len(candidate)+1:
                raise WorkflowError("EXPORT_MISMATCH", "表格列或题数不一致")
            for n, item in enumerate(candidate, 2):
                if [c.value for c in actual[n]] != row_values(item, table, quotes):
                    raise WorkflowError("EXPORT_MISMATCH", "表格内容与候选不一致", scope=item["id"])
                if json.loads(actual.cell(n, 4).value) != item["evidence"]:
                    raise WorkflowError("EXPORT_MISMATCH", "证据数组往返失败")
                expected_main = table[item['evidence'][0]]['url'] if len(item['evidence']) == 1 else None
                actual_main = actual.cell(n,5).hyperlink
                if (actual_main.target if actual_main else None) != expected_main:
                    raise WorkflowError('EXPORT_MISMATCH', '主表原始证据链接不一致')
            detail = check["原始证据链接"]
            if [[detail.cell(fact_header+1+i,j).value for j in range(1,6)] for i in range(len(fact_rows))] != fact_rows:
                raise WorkflowError("EXPORT_MISMATCH", "事实明细往返失败")
            for address, url in link_targets.items():
                if not detail[address].hyperlink or detail[address].hyperlink.target != url:
                    raise WorkflowError("EXPORT_MISMATCH", "来源链接目标不一致")
            if any(c.data_type == "f" for sheet in check for row in sheet for c in row):
                raise WorkflowError("EXPORT_MISMATCH", "业务文本被错误转换为公式")
        finally:
            check.close()
            book.close()
        atomic_bytes(path, data, replace=False)
        return {"question_count": len(candidate), "columns": fields, "fact_count": len(facts),
                "hyperlinks_verified": len(link_targets), "roundtrip": "passed", "status": "待业务人工审核"}

    except WorkflowError:
        raise
    except Exception as exc:
        raise WorkflowError('EXPORT_FAILED', '表格生成或核验失败', scope=type(exc).__name__) from exc
    finally:
        if check is not None:
            check.close()
        book.close()
