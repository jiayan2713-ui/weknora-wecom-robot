"""Independent Mac export; identical archive columns without a Codex runtime."""
import datetime as dt
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def parts(text):
    if not isinstance(text, str) or len(text.encode('utf-16-le')) // 2 <= 32767:
        return []
    return [text[index:index + 15000] for index in range(0, len(text), 15000)]


def limited(text, label):
    return (f'【{label}超过 Excel 单元格上限，完整内容见“采集明细”同序号行的{label}片段列】'
            if parts(text) else text)


def date(value):
    if not value:
        return None
    return dt.datetime.fromisoformat(value.replace('Z', '+00:00')).astimezone(
        dt.timezone(dt.timedelta(hours=8))).replace(tzinfo=None)


def write_excel(records, path):
    fields = [('question', '问题原文'), ('answer', '回答原文'), ('reasoning_process', '推理过程'),
              ('reference_names', '引用资料名称'), ('reference_ids', '引用资料标识JSON'),
              ('usage_json', 'Token用量JSON')]
    data = []
    for record in records:
        refs = record.get('references')
        names = list(dict.fromkeys(name for ref in refs or [] if (name := (
            ref.get('knowledge_title') or ref.get('knowledge_filename') or
            ref.get('knowledge_id') or ref.get('id')))))
        ids = [{key: ref.get(original) for key, original in (
            ('knowledge_title', 'knowledge_title'), ('knowledge_filename', 'knowledge_filename'),
            ('knowledge_id', 'knowledge_id'), ('chunk_id', 'id'), ('knowledge_base_id', 'knowledge_base_id'))}
            for ref in refs] if refs is not None else None
        serialize = lambda value: json.dumps(value, ensure_ascii=False, separators=(',', ':')) if value is not None else None
        data.append({**record, 'reference_names': '\n'.join(names),
                     'reference_ids': serialize(ids), 'usage_json': serialize(record.get('usage'))})
    counts = {key: max((len(parts(row.get(key))) for row in data), default=0) for key, _ in fields}
    headers = ['时间（北京时间）', '用户', '问题', '回答', '状态', '引用资料', '推理过程']
    details = ['主表序号', '记录ID', '会话ID', '请求ID', '问题消息ID', '回答消息ID',
               '回答更新时间（北京时间）', '归属依据', '引用资料标识JSON', 'Token用量JSON', '渠道ID',
               '本地首次采集时间（北京时间）', '本地最后变更时间（北京时间）', '企微聊天ID', '源会话删除时间（北京时间）']
    details += [f'{label}片段{index + 1}' for key, label in fields for index in range(counts[key])]
    main_rows, detail_rows = [], []
    for index, row in enumerate(data, 1):
        main_rows.append([date(row.get('question_at')), row.get('user_id'), limited(row.get('question'), '问题原文'),
                          limited(row.get('answer'), '回答原文'), row['status'], limited(row['reference_names'], '引用资料名称'),
                          limited(row.get('reasoning_process'), '推理过程')])
        detail = [index, row['record_id'], row['session_id'], row.get('request_id'), row.get('user_message_id'),
                  row.get('assistant_message_id'), date(row.get('answer_updated_at')), row.get('attribution'),
                  limited(row['reference_ids'], '引用资料标识JSON'), limited(row['usage_json'], 'Token用量JSON'), row.get('channel_id'),
                  date(row.get('first_collected_at')), date(row.get('last_changed_at')), row.get('chat_id'), date(row.get('session_deleted_at'))]
        for key, _ in fields:
            values = parts(row.get(key))
            detail += values + [None] * (counts[key] - len(values))
        detail_rows.append(detail)
    book = Workbook()
    book.remove(book.active)
    for name, cols, rows, widths in [
        ('问答记录', headers, main_rows, [23, 22, 48, 100, 18, 52, 100]),
        ('采集明细', details, detail_rows, [12, 40, 40, 40, 40, 40, 23, 24, 65, 45, 40, 23, 23, 40, 23])]:
        sheet = book.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet.append([f'机器人 5 · {name}'])
        sheet.append([f'共 {len(records)} 组问答；时间为北京时间。'])
        sheet.append(['一组问答一行；超长原文见采集明细同序号行的片段列。'])
        sheet.append([])
        sheet.append(cols)
        for row in rows:
            sheet.append(row)
        for row in sheet:
            for cell in row:
                # Explicit text prevents formula injection and preserves leading apostrophes.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.font = Font(name='Arial', size=10, color='1F2937')
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if isinstance(cell.value, dt.datetime):
                    cell.number_format = 'yyyy-mm-dd hh:mm:ss'
        sheet['A1'].font = Font(name='Arial', size=14, bold=True)
        for cell in sheet[5]:
            cell.fill = PatternFill('solid', fgColor='25476A')
            cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        for index in range(1, len(cols) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = widths[index - 1] if index <= len(widths) else 50
        for index in range(6, len(rows) + 6):
            sheet.row_dimensions[index].height = 120 if name == '问答记录' else 45
        sheet.freeze_panes = 'A6'
        sheet.auto_filter.ref = f'A5:{get_column_letter(len(cols))}{len(rows) + 5}'
    book.save(path)
    os.chmod(path, 0o600)
