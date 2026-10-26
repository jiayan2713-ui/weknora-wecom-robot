import copy
import json
import tempfile
import unittest
import zipfile
import urllib.error
import struct
import zlib
from pathlib import Path
from unittest.mock import patch

from docflow.common import WorkflowError,artifact,read_json,write_json,digest
from docflow.config import configuration
from docflow.documents import safe_office_copy
from docflow.model import ModelClient
from docflow.parser import parse_document,transcribe
from docflow.storage import Store
from docflow.export import verify_workbook
from tests.fixtures import workflow,start,approve,candidate_files,BODY,SMALL
from tests.test_parser import response


class RegressionTests(unittest.TestCase):
    def test_status_rejects_nonobject(self):
        with tempfile.TemporaryDirectory() as temp:
            w,_=workflow(Path(temp).resolve()/'state')
            with self.assertRaises(WorkflowError):w.execute('status',[])

    def test_auth_failure_in_review_stops_remaining_documents(self):
        from docflow.runner import Workflow
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp).resolve();calls=[]
            for name in ['a.csv','b.csv']:(root/name).write_text('a\nb\n')
            def factory(config):
                def transport(payload):
                    calls.append(1)
                    if '文档转写核对员' in str(payload):
                        raise urllib.error.HTTPError('https://example.invalid',401,'Unauthorized',{},None)
                    return response(BODY)
                return ModelClient(config,transport)
            w=Workflow(root/'state',client_factory=factory)
            result=w.execute('start',{'request_id':'auth','project':'auth fixture','sources':[str(root/'a.csv'),str(root/'b.csv')],'config':SMALL})
            self.assertEqual(result['errors'][0]['code'],'MODEL_HTTP_ERROR')
            self.assertEqual(len(calls),2)
            self.assertEqual(result['documents'][1]['state'],'pending')

    def test_model_record_schema_and_unknown_budget(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp).resolve()
            (root/'broken').mkdir();write_json(root/'broken/call.json',{})
            with self.assertRaises(WorkflowError) as error:ModelClient(configuration(),lambda _:response()).call(root,'broken',[])
            self.assertEqual(error.exception.code,'INTEGRITY_ERROR')
            def interrupted(_):raise SystemExit('simulated process interruption')
            with self.assertRaises(SystemExit):ModelClient(configuration(),interrupted).call(root,'interrupted',[])
            with self.assertRaises(WorkflowError) as error:ModelClient(configuration(),interrupted).call(root,'interrupted',[])
            self.assertEqual(error.exception.code,'MODEL_OUTCOME_UNKNOWN')
            self.assertEqual(len(read_json(root/'interrupted/call.json')['attempts']),1)
            for _ in range(2):
                with self.assertRaises(SystemExit):ModelClient(configuration(),interrupted,retry_terminal=True).call(root,'interrupted',[])
            with self.assertRaises(WorkflowError) as error:ModelClient(configuration(),interrupted,retry_terminal=True).call(root,'interrupted',[])
            self.assertEqual(error.exception.code,'ATTEMPTS_EXHAUSTED')
            self.assertEqual([x['status'] for x in read_json(root/'interrupted/call.json')['attempts']],['unknown']*3)

    def test_review_sentinels_preserve_body(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp).resolve();config=configuration({"model_review": True});prompt=Path(config['parser_prompt']).read_text()
            for i,sentinel in enumerate(['无需修改.','**无需修改。**']):
                replies=iter([response(BODY),response(sentinel)])
                client=ModelClient(config,lambda _:next(replies))
                result=transcribe(client,root/str(i),{'id':'u0001','label':'记录1','text':BODY,'image':None},'.csv',prompt)
                self.assertEqual(result['body'],BODY)

    def test_cache_identity_and_interrupted_checkpoint(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp).resolve();store=Store(root);config=configuration();prompt=Path(config['parser_prompt']).read_text()
            source=root/'test.csv';source.write_text('a\nb\n')
            doc={'name':'test.csv','extension':'.csv','source':artifact(root,source)}
            client=ModelClient(config,lambda p:response('无需修改' if '核对员' in str(p) else BODY))
            parse_document(store,doc,config,prompt,lambda:None,client)
            with self.assertRaises(WorkflowError) as error:parse_document(store,doc,{**config,'temperature':0.9},prompt,lambda:None,client)
            self.assertEqual(error.exception.code,'CONFIG_CHANGED')
            # Keep a trusted unit reference but remove only the final publication pointer.
            doc.pop('cache_ref');directory=store.path(f'parse-cache/{doc["cache_key"]}/{doc["cache_version"]}')
            (directory/'parse-manifest.json').unlink()
            result=store.path(doc['unit_refs']['u0001']['path']);payload=read_json(result);payload['body']='tampered';write_json(result,payload)
            with self.assertRaises(WorkflowError) as error:parse_document(store,doc,config,prompt,lambda:None,client)
            self.assertEqual(error.exception.code,'INTEGRITY_ERROR')

    def test_image_formula_and_malformed_office(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp).resolve();p=root/'image.xlsx';book=openpyxl.Workbook();book.active['A1']='=IMAGE("https://example.invalid/image.png")';book.save(p);book.close()
            with self.assertRaises(WorkflowError) as error:safe_office_copy(p,root)
            self.assertEqual(error.exception.code,'UNSAFE_OFFICE')
            bad=root/'bad.xlsx'
            with zipfile.ZipFile(bad,'w') as z:z.writestr('xl/workbook.xml','<workbook/>')
            w,calls=workflow(root/'state');result=w.execute('start',{'request_id':'bad','project':'bad fixture','sources':[str(bad)],'config':SMALL})
            self.assertEqual(result['documents'][0]['state'],'failed')
            self.assertEqual(result['documents'][0]['issues'][0]['code'],'INVALID_DOCUMENT')
            self.assertEqual(calls,[])

    def test_image_decompression_bomb_is_a_recorded_capacity_error(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp).resolve();path=root/'oversize.png'
            def chunk(kind, data):return struct.pack('>I',len(data))+kind+data+struct.pack('>I',zlib.crc32(kind+data)&0xffffffff)
            path.write_bytes(b'\x89PNG\r\n\x1a\n'+chunk(b'IHDR',struct.pack('>IIBBBBB',20000,10000,8,2,0,0,0))+chunk(b'IDAT',zlib.compress(b'\0'))+chunk(b'IEND',b''))
            w,calls=workflow(root/'state');s=w.execute('start',{'request_id':'bomb','project':'capacity fixture','sources':[str(path)],'config':SMALL})
            self.assertEqual(s['documents'][0]['state'],'failed')
            self.assertEqual(s['documents'][0]['issues'][0]['code'],'CAPACITY_EXCEEDED')
            self.assertEqual(calls,[])

    def test_approve_uses_registered_hash(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp).resolve();w,_=workflow(root/'state');s=start(w,root);original=w._save_request
            changed=False
            def save(req):
                nonlocal changed
                original(req)
                if req['operation']=='approve' and not changed:
                    changed=True;p=Path(s['documents'][0]['working']);p.write_text(p.read_text().replace('10万元','99万元'))
            with patch.object(w,'_save_request',side_effect=save):result=approve(w,s)
            self.assertEqual(result['errors'][0]['code'],'CONTENT_CHANGED')
            self.assertNotIn('generation',result)

    def test_reparse_cannot_reuse_active_approval(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp).resolve();w,_=workflow(root/'state');s=approve(w,start(w,root))
            s=w.execute('resume',{'request_id':'reparse','batch_id':s['batch_id'],'expected_revision':s['revision'],'mode':'reparse','config':{'temperature':0.3}})
            self.assertEqual(s['phase'],'awaiting_review')
            s=w.execute('resume',{'request_id':'regen','batch_id':s['batch_id'],'expected_revision':s['revision'],'mode':'regenerate','codex_model':'fixture'})
            self.assertEqual(s['errors'][0]['code'],'REVIEW_CONFIRMATION_REQUIRED')

    def test_registered_start_recovers_missing_batch(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp).resolve();w,_=workflow(root/'state');save=w._save_batch;failed=False
            def first_failure(batch):
                nonlocal failed
                if not failed:failed=True;raise OSError('simulated first batch write failure')
                return save(batch)
            with patch.object(w,'_save_batch',side_effect=first_failure):
                with self.assertRaises(OSError):start(w,root)
            self.assertEqual(start(w,root)['phase'],'awaiting_review')

    def test_fixed_original_and_replayed_export_integrity(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp).resolve();w,_=workflow(root/'state');s=approve(w,start(w,root));args=candidate_files(s)
            batch=w.store.batch(s['batch_id']);original=w.store.path(batch['documents'][0]['source']['path']);old=original.read_bytes();original.write_bytes(b'changed')
            failed=w.execute('submit',args)
            self.assertEqual(failed['errors'][0]['code'],'INTEGRITY_ERROR')
            original.write_bytes(old)
            s=w.execute('resume',{'request_id':'recover','batch_id':failed['batch_id'],'expected_revision':failed['revision']})
            self.assertEqual(s['phase'],'completed')
            Path(s['artifacts']['original']).write_bytes(b'corrupted generated workbook')
            with self.assertRaises(WorkflowError) as error:w.execute('resume',{'request_id':'recover','batch_id':failed['batch_id'],'expected_revision':failed['revision']})
            self.assertEqual(error.exception.code,'INTEGRITY_ERROR')

    def test_orphan_workbook_is_verified_and_recovered(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp).resolve();w,calls=workflow(root/'state');s=approve(w,start(w,root));args=candidate_files(s);save=w._save_request
            def interrupt(req):
                if req['checkpoints'].get('generated'):raise SystemExit('after workbook, before checkpoint')
                save(req)
            with patch.object(w,'_save_request',side_effect=interrupt):
                with self.assertRaises(SystemExit):w.execute('submit',args)
            s=w.status(s['batch_id']);s=w.execute('resume',{'request_id':'orphan-recovery','batch_id':s['batch_id'],'expected_revision':s['revision']})
            self.assertEqual(s['phase'],'completed',s)
            self.assertEqual(len(calls),2)

    def test_resume_does_not_accept_missing_quote_column_as_legacy(self):
        import openpyxl
        for legacy in [False,True]:
            with self.subTest(legacy=legacy),tempfile.TemporaryDirectory() as temp:
                root=Path(temp).resolve();w,_=workflow(root/'state');s=approve(w,start(w,root));args=candidate_files(s);save=w._save_request
                def interrupt(req):
                    if legacy and req['operation']=='submit':req.pop('xlsx_columns',None)
                    if req['checkpoints'].get('generated'):raise SystemExit('before checkpoint')
                    save(req)
                with patch.object(w,'_save_request',side_effect=interrupt):
                    with self.assertRaises(SystemExit):w.execute('submit',args)
                req=w.store.read('requests/'+args['request_id']+'.json')
                path=w.store.path(f'batches/{s["batch_id"]}/exports/{req["export_id"]}/generated.xlsx')
                book=openpyxl.load_workbook(path)
                if not legacy:book['评测集'].delete_cols(7);book.save(path)
                book.close()
                s=w.status(s['batch_id']);result=w.execute('resume',{'request_id':'resume-schema','batch_id':s['batch_id'],'expected_revision':s['revision']})
                if legacy:self.assertEqual(result['phase'],'completed',result)
                else:self.assertEqual(result['errors'][0]['code'],'EXPORT_MISMATCH',result)

    def test_corrupt_workbook_has_structured_error(self):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'broken.xlsx';path.write_bytes(b'not a workbook')
            with self.assertRaises(WorkflowError) as error:verify_workbook(path,[],[],{}, {})
            self.assertEqual(error.exception.code,'EXPORT_MISMATCH')

    def test_reuse_review_with_excluded_unit(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp).resolve();w,_=workflow(root/'state');source=root/'long.csv';source.write_text('column\n'+'\n'.join(str(i) for i in range(41)))
            args={'request_id':'s1','project':'exclusion fixture','sources':[str(source)],'config':SMALL}
            s=w.execute('start',args)
            s=w.execute('approve',{'request_id':'a1','batch_id':s['batch_id'],'expected_revision':s['revision'],'confirmed':True,
                  'statement':'确认测试范围，排除第二单元','exclude_units':{'d0001':{'u0002':'测试排除'}},'codex_model':'fixture'})
            first=w.store.batch(s['batch_id']);rid=Path(first['review']['path']).parent.name
            second=w.execute('start',{**args,'request_id':'s2'})
            second=w.execute('approve',{'request_id':'a2','batch_id':second['batch_id'],'expected_revision':second['revision'],'confirmed':True,
                 'statement':'明确沿用测试审核','codex_model':'fixture','reuse_review':{'batch_id':first['id'],'review_id':rid},'reuse_map':{'d0001':'d0001'}})
            self.assertEqual(second['phase'],'awaiting_codex',second)

    def test_unrelated_question_change_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp).resolve();w,_=workflow(root/'state');source=root/'source.csv';source.write_text('data\n')
            s=w.execute('start',{'request_id':'s','project':'repair fixture','sources':[str(source)],'config':{**SMALL,'single_step_questions':2}})
            s=approve(w,s);args=candidate_files(s,bad=True)
            candidate=read_json(args['candidate_file']);facts=read_json(args['facts_file']);key=candidate[0]['evidence'][0]
            candidate.append({'id':'QA-002','question':'项目A办理多久？','answer':'办理周期为3个月。','evidence':[key],
                              'required_facts':['办理周期为3个月。'],'category':'知识库内','推理步数':'单步'})
            facts.append({'question_id':'QA-002','fact_index':1,'fact':'办理周期为3个月。','sources':[{'evidence':key,'quote':'办理周期为3个月。'}]})
            write_json(args['candidate_file'],candidate);write_json(args['facts_file'],facts)
            s=w.execute('submit',args)
            candidate[1]['question']='x'*32768
            facts[0]['sources'][0]['quote']=BODY.split('。')[0]+'。'
            target=s['submission'];write_json(target['candidate_file'],candidate);write_json(target['facts_file'],facts)
            s=w.execute('submit',{**args,**target,'expected_revision':s['revision']})
            self.assertIn('UNRELATED_REPAIR',[e['code'] for e in s['validation']['errors']])
            candidate[1]['question']='另一种正常长度的无关改写'
            target=s['submission'];write_json(target['candidate_file'],candidate);write_json(target['facts_file'],facts)
            s=w.execute('submit',{**args,**target,'expected_revision':s['revision']})
            self.assertNotEqual(s['phase'],'completed')
            self.assertIn('UNRELATED_REPAIR',[e['code'] for e in s['validation']['errors']])

    def test_submit_uses_registered_input_hashes(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp).resolve();w,_=workflow(root/'state');s=approve(w,start(w,root));args=candidate_files(s);save=w._save_request
            changed=False
            def mutate(req):
                nonlocal changed
                save(req)
                if req['operation']=='submit' and not changed:
                    changed=True;data=read_json(args['candidate_file']);data[0]['answer']='注册后变更';write_json(args['candidate_file'],data)
            with patch.object(w,'_save_request',side_effect=mutate):result=w.execute('submit',args)
            self.assertEqual(result['errors'][0]['code'],'CONTENT_CHANGED')

    def test_default_fifty_rows_and_native_hyperlinks(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp).resolve();w,_=workflow(root/'state');source=root/'source.csv';source.write_text('data\n')
            s=w.execute('start',{'request_id':'default50','project':'字段与导出测试，不评判题目语义','sources':[str(source)],'urls':{'source.csv':'https://example.invalid/source'}})
            s=approve(w,s);table=read_json(s['generation']['evidence_table']);key=next(iter(table));candidate=[];facts=[]
            for i in range(1,51):
                category='知识库内' if i<=40 else ('知识库外' if i<=45 else '模糊地带')
                item={'id':f'QA-{i:03}','question':f'字段测试{i}（仅测试结构）','answer':BODY,'evidence':[] if category=='知识库外' else [key],
                      'required_facts':[BODY] if category!='知识库外' else [],'category':category,
                      '推理步数':('单步' if i<=32 else '多步') if i<=40 else '不适用'}
                candidate.append(item)
                if category!='知识库外':facts.append({'question_id':item['id'],'fact_index':1,'fact':BODY,'sources':[{'evidence':key,'quote':BODY}]})
            target=s['submission'];write_json(target['candidate_file'],candidate);write_json(target['facts_file'],facts)
            result=w.execute('submit',{**target,'batch_id':s['batch_id'],'expected_revision':s['revision'],'generation_id':s['generation']['id']})
            self.assertEqual(result['phase'],'completed',result)
            book=openpyxl.load_workbook(result['artifacts']['original'])
            self.assertEqual(book['评测集'].max_row,51)
            self.assertEqual(book['评测集'].max_column,9)
            self.assertIn(BODY,book['评测集'].cell(47,7).value)
            self.assertEqual(book['评测集'].cell(47,6).value,'1. '+BODY)
            self.assertEqual(book['评测集'].cell(42,6).value,'[]')
            self.assertEqual(read_json(result['artifacts']['manifest'])['checks']['fact_count'],45)
            self.assertEqual(sum(c.hyperlink is not None for sheet in book for row in sheet for c in row),90)
            book.close()
