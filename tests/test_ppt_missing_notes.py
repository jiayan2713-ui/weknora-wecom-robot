import tempfile,unittest,json
from pathlib import Path
from unittest.mock import patch
from PIL import Image
from pptx import Presentation
from pypdf import PdfWriter
from docflow.config import configuration
from docflow.documents import prepare

class PptNotesTests(unittest.TestCase):
 def test_real_package_without_notes_text_frame_preserves_other_notes(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp).resolve();source=root/'notes.pptx';prs=Presentation()
   first=prs.slides.add_slide(prs.slide_layouts[6]);first.notes_slide.notes_text_frame.text='需要保留的备注：123万元'
   missing=prs.slides.add_slide(prs.slide_layouts[6]);notes=missing.notes_slide
   for shape in list(notes.shapes):shape.element.getparent().remove(shape.element)
   self.assertTrue(missing.has_notes_slide);self.assertIsNone(notes.notes_text_frame)
   last=prs.slides.add_slide(prs.slide_layouts[6]);self.assertFalse(last.has_notes_slide)
   prs.save(source)
   def fake_run(command,timeout,**kwargs):
    if '--convert-to' in command:
     out=Path(command[command.index('--outdir')+1]);writer=PdfWriter()
     for _ in range(3):writer.add_blank_page(width=100,height=100)
     writer.write(out/(Path(command[-1]).stem+'.pdf'))
    else:Image.new('RGB',(20,20),'white').save(command[-1]+'.png')
   with patch('docflow.documents.run',side_effect=fake_run):units=prepare(source,root/'prepared',configuration())
   data=json.loads((root/'prepared/prepared.json').read_text())
   self.assertEqual([s['notes'] for s in data['native']['slides']],['需要保留的备注：123万元','',''])
   self.assertEqual(len(units),3)
   self.assertIn('需要保留的备注：123万元',units[0]['text'])
   self.assertTrue(Presentation(source).slides[1].has_notes_slide)

if __name__=='__main__':unittest.main()
